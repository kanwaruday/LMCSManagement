#!/usr/bin/env python3
"""Agency service data -> the parts catalog. The vehicle agencies' service centres fill the Excel template
(transport/game-data/Service centre data template (agencies).xlsx: Our vehicles, Service history, Maintenance schedule, Wear parts,
Service packages, Running norms, Contacts). This reads one or more filled files and pushes them to the portal (POST
transportcatalogpush, same secret and helper as capture.py) as the parts catalog the garage reads:
catalog.models[model].parts[code] = {lifeKm, lifeMonths, price, labour, spec, components, asOf}.

  /usr/local/bin/python3 transport/capture/agency_import.py FILE.xlsx [FILE2.xlsx ...]          # dry run: a summary, nothing sent
  /usr/local/bin/python3 transport/capture/agency_import.py FILE.xlsx [FILE2.xlsx ...] --push   # send it (exit non-zero if refused)
  ... --as-of 2026-10-20                                                                         # the date the agency's data is from
  /usr/local/bin/python3 transport/capture/test_agency_import.py                                 # tests (framework 3.13: has openpyxl)

Reading rules: the template's example rows (recognised by their example values, never by font or colour; each skip is noted) and
blank rows are skipped. Dates are DD-MM-YYYY, money is rupees with or without ₹ and
commas. A row with text where a number belongs (km, months, price, labour, odometer, amount) or a date that is not a date is NOT
imported; it is listed under "not imported" so the agency can fix it. The Model column of an item block may be filled on the first
row only (it carries down). Maintenance schedule: the severe/hill-duty interval wins over the normal one when present, per field.
Wear parts give the replacement life, price and labour, and win over the schedule for the same model and part. Several files for
one model merge, the newest file (by its modified time, or --as-of YYYY-MM-DD, kept as asOf) winning per part. Every vehicle gets
`catalogModel`, the catalog model its model text matches (exactly, else by one unique prefix, else it is reported). Registrations are
stored without any non-alphanumeric character; a vehicle or history row whose registration cannot be found is reported, not pushed.

The template's codes stay the catalog keys; each part also carries `components`, the Task 1 ids (OFFICE_COMPONENTS in
apps-script/office-config.gs) the garage matches service bills against. Contacts (advisors' phone numbers) and the packages and
norms sheets are not pushed. Money is business data for the private Sheet only. Nothing here computes speed.
"""
import argparse, datetime as dt, re, sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import capture  # post(): the secret from ~/.lmcs/transport_push_secret, curl --data-binary @-, exits non-zero if refused

# Task 1's ids, OFFICE_COMPONENTS in apps-script/office-config.gs (test_agency_import.py fails if this copy drifts)
OFFICE_IDS = frozenset("""engine_oil oil_filter fuel_filter air_filter coolant belts brake_front brake_rear brake_fluid clutch gearbox_oil
axle_grease suspension steering wheel_bearing battery lights_wiring starter_alternator tyre_fl tyre_fr tyre_rl tyre_rr tyre_spare
puncture body_paint seats_windows door washing_greasing other""".split())
# template code -> Task 1 ids where they differ (Ruling 2). Any other code maps to itself when that is a Task 1 id; the rest stay
# catalog-only (axle_oil, ps_fluid, def, injector, prop_shaft, radiator in today's template).
CODE_MAP = {"tyre": ["tyre_fl", "tyre_fr", "tyre_rl", "tyre_rr", "tyre_spare"], "leaf_spring": ["suspension"], "brake_drum": ["brake_front", "brake_rear"]}

SHEETS = {
    "vehicles": ("Our vehicles", {"chassis": "chassis", "reg": "registration", "model": "model", "variant": "variant", "engine": "engine",
                                  "emission": "emission", "year": "year", "warrantyDate": "warranty valid till (date", "warrantyKm": "warranty valid till (km",
                                  "recalls": "open recalls"}),
    "history": ("Service history", {"chassis": "chassis", "jobCard": "job card", "date": "date", "odometer": "odometer", "work": "work done",
                                    "part": "part name", "partNo": "part number", "qty": "quantity", "partAmount": "part amount", "labour": "labour"}),
    "schedule": ("Maintenance schedule", {"code": "code", "model": "model", "nKm": "normal: km", "nMo": "normal: months",
                                          "sKm": "severe / hill duty: km", "sMo": "severe / hill duty: months", "spec": "grade"}),
    "wear": ("Wear parts", {"code": "code", "model": "model", "km": "expected life on", "mo": "expected life: m", "price": "price",
                            "labour": "labour", "notes": "notes"}),
    "contacts": ("Contacts", {"agency": "agency", "advisor": "service advisor", "phone": "phone", "place": "workshop (place", "nearest": "nearest",
                              "timings": "workshop timings", "pickup": "pick-up", "booking": "how to book"}),
}


def _same(v, want):
    try:
        return number(v) == number(want)
    except ValueError:
        return text(v) is not None and text(v).lower() == str(want).lower()


def _sig(**want):
    return lambda r: all(_same(r[k], w) for k, w in want.items())


# The template's example rows, known by their VALUES (product names and placeholder marks only): an agency's re-save loses the grey
# italic, and data typed over an example row keeps it, so the font tells nothing. The "…" marks a truncated chassis or part number.
EXAMPLES = {
    "vehicles": lambda r: "…" in str(r["chassis"] or ""),
    "history": lambda r: "…" in str(r["chassis"] or ""),
    "schedule": _sig(code="engine_oil", model="LP 410/36", nKm=20000, nMo=12, sKm=10000, sMo=6, spec="15W-40 CI-4"),
    "wear": _sig(code="brake_front", model="LP 410/36", km=35000, mo=18, price=2400, labour=900),
    "contacts": lambda r: "xx" in str(r["phone"] or "").lower(),
}
HISTORY_KEY = ("chassis", "jobCard", "date", "part", "partNo", "qty")


def components_for(code):
    c = str(code or "").strip().lower()
    return list(CODE_MAP[c]) if c in CODE_MAP else [c] if c in OFFICE_IDS else []


def text(v):
    s = " ".join(str(v).split()) if v is not None else ""
    return s or None


def number(v):
    """None for an empty cell, an int/float for a number (₹, commas, a trailing km/months allowed); ValueError for anything else."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError(f"{v!r} is not a number")
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    m = re.fullmatch(r"₹?\s*(\d[\d,]*(?:\.\d+)?)\s*(?:km|kms|months?|mo)?", str(v).strip().lower())
    if not m:
        raise ValueError(f"{str(v).strip()!r} is not a number")
    f = float(m[1].replace(",", ""))
    return int(f) if f.is_integer() else f


def date(v):
    """ISO yyyy-mm-dd from a DD-MM-YYYY text (or an Excel date); None for an empty cell; ValueError otherwise."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", str(v).strip())
    try:
        return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
    except (TypeError, ValueError):
        raise ValueError(f"{str(v).strip()!r} is not a DD-MM-YYYY date") from None


def ident(v):  # chassis numbers: no spaces, upper case
    s = re.sub(r"\s+", "", str(v)).upper() if v is not None else ""
    return s or None


def reg_id(v):
    """Registration number as the rest of the repo keeps it (officeNormReg_, tupNormReg_): upper case, every non-alphanumeric dropped."""
    s = re.sub(r"[^A-Z0-9]", "", str(v).upper()) if v is not None else ""
    return s or None


def money(v):
    """number(), but a dash, n/a or nil means empty (an agency's way of saying 'no price'). Km and months stay strict."""
    if isinstance(v, str) and re.fullmatch(r"[\s\-–—.]*|n/?a|nil|none", v.strip().lower()):
        return None
    return number(v)


def table(wb, key, data):
    """Yields (row number, {field: raw value}, kind) for the rows of a sheet below its header row (found by header text, so a changed
    intro line or an inserted column does not matter). kind: 'example' (the template's example values), 'data' (something in a `data`
    field), 'model' (only a model typed; the caller must still carry it down). Rows with nothing in them are not yielded."""
    name, fields = SHEETS[key]
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    for hr in range(1, 8):
        heads = [" ".join(str(c.value or "").lower().split()) for c in ws[hr]]
        cols = {f: next((i for i, h in enumerate(heads) if h.startswith(p)), None) for f, p in fields.items()}
        if all(i is not None for i in cols.values()):
            break
    else:
        raise ValueError(f"sheet '{name}': the column headers were not found (do not rename the columns)")
    for row in ws.iter_rows(min_row=hr + 1):
        vals = {f: row[i].value if i < len(row) else None for f, i in cols.items()}
        if EXAMPLES[key](vals):
            yield row[0].row, vals, "example"
        elif any(text(vals[f]) for f in data):
            yield row[0].row, vals, "data"
        elif text(vals.get("model")):
            yield row[0].row, vals, "model"


def parse_agency(path, as_of=None):
    """One filled workbook -> {models, vehicles, history, contacts, problems, notes}. models: [{model, parts: {code: {lifeKm,
    lifeMonths, price, labour, spec, components, asOf}}}]; problems: one line per row left out and why ("file: sheet row n: ...");
    notes: the example rows that were skipped. as_of: 'YYYY-MM-DD' for the parts' asOf (default: the file's modified time)."""
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    as_of = f"{as_of}T00:00:00" if as_of else dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%dT%H:%M:%S")
    out = {"models": [], "vehicles": [], "history": [], "contacts": [], "problems": [], "notes": []}
    bad = lambda sheet, n, why: out["problems"].append(f"{path.name}: {sheet} row {n}: {why}")

    def rows(key, data):
        try:
            for n, r, kind in table(wb, key, data):
                if kind == "example":
                    out["notes"].append(f"{path.name}: {SHEETS[key][0]} row {n} is the template's example row and was skipped "
                                        f"(if it holds your own data, change it so it differs from the example)")
                else:
                    yield n, r, kind
        except ValueError as e:
            out["problems"].append(f"{path.name}: {e}")

    for n, r, _ in rows("vehicles", list(SHEETS["vehicles"][1])):
        try:
            v = {"chassis": ident(r["chassis"]), "reg": reg_id(r["reg"]), "model": text(r["model"]), "variant": text(r["variant"]),
                 "engine": text(r["engine"]), "emission": text(r["emission"]), "year": number(r["year"]), "warrantyDate": date(r["warrantyDate"]),
                 "warrantyKm": number(r["warrantyKm"]), "recalls": text(r["recalls"]), "asOf": as_of, "file": path.name, "row": n}
        except ValueError as e:
            bad("Our vehicles", n, e)
            continue
        if not v["chassis"]:
            bad("Our vehicles", n, "no chassis number")
        elif not v["reg"]:
            bad("Our vehicles", n, f"{v['chassis']} has no registration number, so it cannot be shown for any bus (not imported)")
        else:
            out["vehicles"].append(v)
    reg_of = {v["chassis"]: v["reg"] for v in out["vehicles"]}

    for n, r, _ in rows("history", list(SHEETS["history"][1])):
        try:
            h = {"chassis": ident(r["chassis"]), "jobCard": text(r["jobCard"]), "date": date(r["date"]), "odometer": number(r["odometer"]),
                 "work": text(r["work"]), "part": text(r["part"]), "partNo": text(r["partNo"]), "qty": text(r["qty"]),
                 "partAmount": money(r["partAmount"]), "labour": money(r["labour"]), "file": path.name, "row": n}
        except ValueError as e:
            bad("Service history", n, e)
            continue
        if not (h["chassis"] and h["date"]):
            bad("Service history", n, "needs a chassis number and a date")
        else:
            h["reg"] = reg_of.get(h["chassis"], "")
            out["history"].append(h)

    models = {}
    for key, sheet, data, extract in (
            ("schedule", "Maintenance schedule", ["nKm", "nMo", "sKm", "sMo", "spec"],
             lambda r: {"lifeKm": first(number(r["sKm"]), number(r["nKm"])), "lifeMonths": first(number(r["sMo"]), number(r["nMo"])), "spec": text(r["spec"])}),
            ("wear", "Wear parts", ["km", "mo", "price", "labour", "notes"],
             lambda r: {"lifeKm": number(r["km"]), "lifeMonths": number(r["mo"]), "price": money(r["price"]), "labour": money(r["labour"]), "spec": text(r["notes"])})):
        model, seen = None, set()  # ponytail: the Model column may be filled on the first row of a block only; it carries down (even from a row with no data)
        for n, r, kind in rows(key, data):
            if text(r["model"]):
                model = " ".join(str(r["model"]).split())
            if kind != "data":
                continue
            code = text(r["code"]) and text(r["code"]).lower()
            try:
                vals = extract(r)
            except ValueError as e:
                bad(sheet, n, f"{code}: {e}")
                continue
            if not (code and model):
                bad(sheet, n, "needs a code and a model")
            elif (model, code) in seen:  # a second block whose Model was not typed would otherwise overwrite the first
                bad(sheet, n, f"{code} is listed twice for {model!r}; type the Model on every block (this row was not imported)")
            else:
                seen.add((model, code))
                p = models.setdefault(model, {}).setdefault(code, {"lifeKm": None, "lifeMonths": None, "price": None, "labour": None, "spec": None})
                p.update({k: v for k, v in vals.items() if v is not None})  # the wear sheet is read last: its life, price and labour win
    out["models"] = [{"model": m, "parts": {c: dict(p, components=components_for(c), asOf=as_of) for c, p in parts.items()}} for m, parts in models.items()]

    for n, r, _ in rows("contacts", list(SHEETS["contacts"][1])):
        out["contacts"].append({k: text(v) for k, v in r.items()})
    return out


def first(*vals):
    return next((v for v in vals if v is not None), None)


def catalog_key(model, keys):
    """The catalog model a vehicle's model text means: an exact match (any case), else the one catalog model that is a prefix of it (or
    that it is a prefix of); (key, '') or ('', why) when there is none or more than one."""
    m = (model or "").lower()
    if not m:
        return "", "it has no model"
    exact = [k for k in keys if k.lower() == m]
    if exact:
        return exact[0], ""
    near = [k for k in keys if m.startswith(k.lower()) or k.lower().startswith(m)]
    if len(near) == 1:
        return near[0], ""
    return "", f"matches no catalog model" if not near else f"matches more than one catalog model ({', '.join(sorted(near))})"


def merge(parsed):
    """Several parse_agency results -> one {models, vehicles, history, contacts, problems, notes}. Per model and part, and per vehicle,
    the newest asOf wins (a tie: the later file). History rows are listed once per job card line (the later file wins). Vehicles get
    their catalogModel; history rows whose chassis is on no vehicle sheet are reported and left out."""
    models, vehicles, history = {}, {}, {}
    for p in parsed:
        for m in p["models"]:
            mine = models.setdefault(m["model"], {})
            for code, part in m["parts"].items():
                if code not in mine or part["asOf"] >= mine[code]["asOf"]:
                    mine[code] = part
        for v in p["vehicles"]:
            if v["chassis"] not in vehicles or v["asOf"] >= vehicles[v["chassis"]]["asOf"]:
                vehicles[v["chassis"]] = dict(v)
        history.update({tuple(h.get(k) for k in HISTORY_KEY): dict(h) for h in p["history"]})
    problems = [x for p in parsed for x in p["problems"]]
    for v in vehicles.values():
        v["catalogModel"], why = catalog_key(v["model"], list(models))
        if why:
            problems.append(f"{v['file']}: Our vehicles row {v['row']}: {v['reg']} (model {v['model']!r}) {why}; the garage cannot price it")
    reg_of = {c: v["reg"] for c, v in vehicles.items()}
    kept = []
    for h in history.values():
        h["reg"] = h.get("reg") or reg_of.get(h["chassis"], "")
        if h["reg"]:
            kept.append(h)
        else:
            problems.append(f"{h['file']}: Service history row {h['row']}: chassis {h['chassis']} is not on any 'Our vehicles' sheet, so the row "
                            f"cannot be shown for a bus (not imported)")
    return {"models": [{"model": m, "parts": parts} for m, parts in models.items()], "vehicles": list(vehicles.values()), "history": kept,
            "contacts": [c for p in parsed for c in p["contacts"]], "problems": problems, "notes": [x for p in parsed for x in p["notes"]]}


def summary(label, p):
    parts = sum(len(m["parts"]) for m in p["models"])
    return (f"{label}: {len(p['models'])} models ({parts} parts), {len(p['vehicles'])} vehicles, {len(p['history'])} history rows, "
            f"{len(p['contacts'])} contacts")


def as_of_date(v):
    try:
        return dt.date.fromisoformat(v).isoformat()
    except ValueError:
        raise argparse.ArgumentTypeError(f"{v!r} is not a YYYY-MM-DD date") from None


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("files", nargs="+", help="filled agency workbook(s) (.xlsx)")
    ap.add_argument("--push", action="store_true", help="send it to the portal (default: a dry run that prints a summary)")
    ap.add_argument("--as-of", type=as_of_date, metavar="YYYY-MM-DD", help="the date the data is from (default: each file's modified time)")
    a = ap.parse_args(argv)
    parsed = [parse_agency(f, a.as_of) for f in a.files]
    for f, p in zip(a.files, parsed):
        print(summary(Path(f).name, p))
    m = merge(parsed)
    print(summary("total", m))
    for line in m["notes"]:
        print("  note:", line)
    if m["problems"]:
        print(f"{len(m['problems'])} problems (ask the agency to fix the rows):")
        for line in m["problems"]:
            print("  -", line)
    if not a.push:
        print("dry run: nothing sent (add --push to send it)")
        return 0
    strip = lambda rows, keys: [{k: v for k, v in r.items() if k not in keys} for r in rows]
    res = capture.post("transportcatalogpush", {"models": m["models"], "vehicles": strip(m["vehicles"], ("file", "row", "asOf")),
                                                "history": strip(m["history"], ("file", "row"))})
    print("pushed:", {k: v for k, v in res.items() if k != "success"})
    return 0


if __name__ == "__main__":
    sys.exit(main())

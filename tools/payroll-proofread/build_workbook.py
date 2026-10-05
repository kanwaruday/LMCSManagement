#!/usr/bin/env python3
"""Builds the Payroll Proofreading Workbook: every monthly-salary formula as live Excel
formulas (one row per employee), with the inputs on tabs named exactly like the Google Sheet
tabs they are copied from, so a fresh export can simply be pasted over them.

The formulas mirror apps-script/payroll.gs (payCompute_, payEsiCovered_, payEmployedDays_,
payLoanSchedules_, payLedgerBalances_) and apps-script/payroll-calc.gs (payCalc_) one for one;
tools/payroll-proofread/verify.py recalculates the workbook in LibreOffice against the real
backend on randomised fake data to prove they agree.

No real data lives in this repo -- build.py loads real exports from paths you pass on the CLI.
"""
import datetime as dt
import re
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# Tab layouts (header text must match the Google Sheets exactly -- see PAY_TABS in payroll.gs).
HEADERS = {
    'PayRoll Constants': ['Setting', 'Value'],
    'PayRoll Rates': ['School', 'DA %', 'ADA %', 'EPF Enrolled by Default', 'ESI Enrolled by Default'],
    'Salary Master': ['Employee Code', 'Name', 'Entity', 'Designation', 'Date of Joining', 'Effective From', 'Full Basic',
                      'EPF Member (Y/N)', 'RRF Member (Y/N)', 'Staff-Child Tuition', 'Pay Mode', 'Bank Account No', 'IFSC', 'UAN', 'PAN',
                      'Status (Active/Left)', 'Remarks', 'Last Working Day', 'Monthly TDS (₹)'],
    'Monthly Inputs': ['Month', 'Employee Code', 'Paid Days', 'CL Days Encashed', 'Hold for F&F (Y/N)', 'Other Earnings',
                       'Other Earnings Note', 'Other Deductions', 'Other Deductions Note', 'Updated By', 'Updated At',
                       'Hold (F&F/Grievance)', 'Release Held (₹)', 'Skip Loan Recovery (Y/N)', 'Skip Reason'],
    'Adjustments': ['ID', 'Month', 'Employee Code', 'Type', 'Direction', 'Amount', 'Note', 'Added By', 'Added At'],
    'Ledger': ['Date', 'Month', 'Employee Code', 'Account (RRF/Security/Loan/Held Salary)',
               'Type (Opening/Deduction/Payout/Loan Issued/Loan Repaid/Adjustment)', 'Amount', 'Reference', 'Notes', 'Entered By', 'Entered At'],
    'Loans': ['Loan ID', 'Employee Code', 'Type (Salary Advance/RRF Loan)', 'Given On', 'Amount', 'Monthly Recovery', 'Start Month',
              'Status (Active/Closed)', 'Notes', 'Tally Ledger Name', 'Added By', 'Added At'],
    'Payroll Register': ['Month', 'Employee Code', 'Name', 'Entity', 'Designation', 'Paid Days', 'Month Basic', 'ADA', 'DA',
                         'CL Encashment', 'Tuition (A)', 'Other Earnings', 'Gross', 'EPF Employee', 'ESI Employee', 'RRF Rate', 'RRF',
                         'Tuition (D)', 'Other Deductions', 'Total Deductions', 'Net Pay', 'Bank Payable', 'Held for F&F', 'EPF Employer',
                         'ESI Employer', 'CTI', 'Gratuity Provision', 'Locked At', 'Locked By', 'Withheld (Grievance)', 'Released Held',
                         'Loan Recovery', 'TDS'],
}
# Rows the helper formulas / lookup ranges cover on each tab (extend by dragging the helper cells down).
LIMIT = {'PayRoll Constants': 200, 'PayRoll Rates': 60, 'Salary Master': 3000, 'Monthly Inputs': 12000, 'Adjustments': 6000,
         'Ledger': 15000, 'Loans': 1000, 'Payroll Register': 15000}
# Helper column (letter) holding the lookup key/date on each tab.
HELP = {'Salary Master': 'U', 'Monthly Inputs': 'Q', 'Adjustments': 'K', 'Ledger': 'K', 'Loans': 'M', 'Payroll Register': 'AI', 'PayRoll Rates': 'F'}

BLUE = Font(color='0000CC')
BOLD = Font(bold=True)
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(bold=True, color='FFFFFF')
GROUPS = {  # fill colour per column group on the calc sheet
    'who': 'DDEBF7', 'master': 'E2EFDA', 'input': 'FFF2CC', 'calc': 'FFFFFF', 'result': 'FCE4D6', 'check': 'EDEDED',
}
thin = Side(style='thin', color='BFBFBF')


def rng(tab, col):
    return "'%s'!$%s$1:$%s$%d" % (tab, col, col, LIMIT[tab])


def num(x):  # payNum_: numbers as-is, text parsed, anything else 0
    return 'IF(ISNUMBER(%s),%s,IFERROR(VALUE(SUBSTITUTE(%s,",","")),0))' % (x, x, x)


def yes(x):  # payYes_: /^y/i
    return 'IF(LEFT(LOWER(TRIM(%s)),1)="y",1,0)' % x


def rnd(x):  # payRound_: Excel's ROUND with float noise removed first
    return 'ROUND(ROUND(%s,6),0)' % x


def up(x):  # round UP to the next rupee (EPF and ESI -- confirmed with Pawan Sir 2026-10-01)
    return 'ROUNDUP(ROUND(%s,6),0)' % x


# ---------------------------------------------------------------------------------------------
# Calc-sheet column definitions.  Formulas use {key} for another column of the SAME row.
# W(...) = "only for employees who are on this month's payroll".
# ---------------------------------------------------------------------------------------------
SM = lambda c: "INDEX(%s,{smrow})" % rng('Salary Master', c)
RT = lambda c: "INDEX(%s,{rateRow})" % rng('PayRoll Rates', c)
MI = lambda c: "INDEX(%s,{miRow})" % rng('Monthly Inputs', c)
W = lambda e: 'IF({incl}="Yes",%s,"")' % e

COLS = []


def col(key, header, formula, group, fmt=None, desc='', src='', code='', width=12, source=None):
    """source = (tab, column-letter, header) for cells whose value is fetched from an input tab."""
    COLS.append(dict(key=key, header=header, formula=formula, group=group, fmt=fmt, desc=desc, src=src, code=code, width=width, source=source))


# --- who / which Salary Master row -----------------------------------------------------------
col('code', 'Employee Code', None, 'who', None, 'Typed in. One row per code that has ever been on Salary Master -- add a row by copying one down and typing a code.',
    'You type it (copied from Salary Master › Employee Code).', 'payCompute_: Object.keys(master)', 16)
col('smrow', 'Salary Master row used', '=IFERROR(MATCH({code}&"|"&TEXT(_xlfn.MAXIFS(%s,%s,{code},%s,"<="&ME),"yyyymmdd"),%s,0),"")' %
    (rng('Salary Master', 'V'), rng('Salary Master', 'Y'), rng('Salary Master', 'V'), rng('Salary Master', 'U')), 'who', '0',
    'Sheet row number on the Salary Master tab of the row in force this month: the latest row for this code whose Effective From is on or before the month end.',
    'Salary Master › Employee Code + Effective From (helper cols U,V,Y)', 'payCompute_: master[code] = latest row with Effective From <= month end', 11)
col('name', 'Name', '=IF({smrow}="","",TRIM(%s))' % SM('B'), 'master', None, 'Name from the Salary Master row in force.', 'Salary Master › Name (col B) at the row shown in "Salary Master row used"', 'payCompute_: r[\'Name\']', 24, ('Salary Master', 'B', 'Name'))
col('entRaw', 'School (as written)', '=IF({smrow}="","",%s)' % SM('C'), 'master', None, 'Entity text as typed on Salary Master.', 'Salary Master › Entity (col C)', 'payCompute_: r[\'Entity\']', 10, ('Salary Master', 'C', 'Entity'))
col('ent', 'School key', '=IF({smrow}="","",SUBSTITUTE(UPPER(%s)," ",""))' % SM('C'), 'master', None, 'Entity with spaces removed and upper-cased so "LMS 1" and "LMS1" match.', 'derived from School (as written)', 'payEntity_()', 9)
col('desig', 'Designation', '=IF({smrow}="","",TRIM(%s))' % SM('D'), 'master', None, 'Designation from Salary Master.', 'Salary Master › Designation (col D)', 'payCompute_', 22, ('Salary Master', 'D', 'Designation'))
col('doj', 'Date of Joining', '=IF({smrow}="","",%s)' % SM('W'), 'master', 'dd-mmm-yyyy', 'Joining date (helper col W converts text dates).', 'Salary Master › Date of Joining (col E)', 'payCompute_: payDate_(r[\'Date of Joining\'])', 12, ('Salary Master', 'E', 'Date of Joining'))
col('lwd', 'Last Working Day', '=IF({smrow}="","",%s)' % SM('X'), 'master', 'dd-mmm-yyyy', 'Last working day if set (blank = still working).', 'Salary Master › Last Working Day (col R)', 'payCompute_: payDate_(r[\'Last Working Day\'])', 12, ('Salary Master', 'R', 'Last Working Day'))
col('status', 'Status', '=IF({smrow}="","",%s)' % SM('P'), 'master', None, 'Status (Active/Left). Only used when there is no Last Working Day.', 'Salary Master › Status (Active/Left) (col P)', 'payCompute_', 9, ('Salary Master', 'P', 'Status (Active/Left)'))
col('rateRow', 'PayRoll Rates row', '=IF({ent}="","",IFERROR(MATCH({ent},%s,0),""))' % rng('PayRoll Rates', 'F'), 'master', '0',
    'Row on the PayRoll Rates tab for this school (helper col F = School without spaces, upper-case).', 'PayRoll Rates › School (col A)', 'payRates_ / rates[entity]', 9)
col('active', 'On payroll this month?', '=IF({smrow}="","",IF({lwd}<>"",{lwd}>=MS,LEFT(LOWER(TRIM({status})),4)<>"left"))', 'master', None,
    'TRUE unless the person left before this month: if there is a Last Working Day it must be on/after the 1st of the month; otherwise Status must not start with "Left".',
    'Last Working Day / Status', 'payCompute_: lwd ? lwd < start : /^left/i.test(status) -> skipped', 10)
col('incl', 'Included?', '=IF(AND({smrow}<>"",{rateRow}<>"",{doj}<>"",{active}=TRUE),"Yes","No")', 'master', None,
    '"Yes" when the person is paid this month. "No" rows stay blank below (reason in the next column).', 'derived', 'payCompute_ (skips left staff, missing rates, missing joining date)', 9)
col('why', 'If not included, why', '=IF({incl}="Yes","",IF({smrow}="","No Salary Master row effective by month end",IF({active}=FALSE,"Left before this month",IF({rateRow}="","No PayRoll Rates row for this school","No Date of Joining"))))',
    'master', None, 'Reason the row is empty.', 'derived', 'payCompute_ warnings', 26)

# --- Salary Master facts ---------------------------------------------------------------------
col('fullBasic', 'Full Basic', W(num(SM('G'))), 'master', '#,##0', 'Full-month basic from Salary Master.', 'Salary Master › Full Basic (col G)', 'payNum_(r[\'Full Basic\'])', 11, ('Salary Master', 'G', 'Full Basic'))
col('epfMem', 'EPF member (1/0)', W(yes(SM('H'))), 'master', '0', '1 if "EPF Member (Y/N)" starts with Y.', 'Salary Master › EPF Member (Y/N) (col H)', 'payYes_', 8, ('Salary Master', 'H', 'EPF Member (Y/N)'))
col('rrfMem', 'RRF member (1/0)', W(yes(SM('I'))), 'master', '0', '1 if "RRF Member (Y/N)" starts with Y.', 'Salary Master › RRF Member (Y/N) (col I)', 'payYes_', 8, ('Salary Master', 'I', 'RRF Member (Y/N)'))
col('tuition', 'Staff-child tuition', W(num(SM('J'))), 'master', '#,##0', 'Monthly staff-child tuition value (added to gross, deducted again -- nets to zero).', 'Salary Master › Staff-Child Tuition (col J)', 'payNum_', 10, ('Salary Master', 'J', 'Staff-Child Tuition'))
col('tdsMaster', 'Monthly TDS (Salary Master)', W(num(SM('S'))), 'master', '#,##0', 'Standing monthly income-tax deduction.', 'Salary Master › Monthly TDS (₹) (col S)', "payNum_(r['Monthly TDS (₹)'])", 10, ('Salary Master', 'S', 'Monthly TDS (₹)'))
col('daPct', 'DA %', W(num(RT('B')) + '/100'), 'master', '0.00%', "School's DA percentage.", 'PayRoll Rates › DA % (col B) at the row shown', 'payRates_: da', 8, ('PayRoll Rates', 'B', 'DA %'))
col('adaPct', 'ADA %', W(num(RT('C')) + '/100'), 'master', '0.00%', "School's ADA percentage.", 'PayRoll Rates › ADA % (col C)', 'payRates_: ada', 8, ('PayRoll Rates', 'C', 'ADA %'))
col('epfReg', 'School EPF-registered (1/0)', W(yes(RT('D'))), 'master', '0', 'Schools marked N here pay no EPF (LMS 6).', 'PayRoll Rates › EPF Enrolled by Default (col D)', 'payRates_: epfRegistered', 9, ('PayRoll Rates', 'D', 'EPF Enrolled by Default'))
col('esiReg', 'School ESI-registered (1/0)', W(yes(RT('E'))), 'master', '0', 'Schools marked N here pay no ESI.', 'PayRoll Rates › ESI Enrolled by Default (col E)', 'payRates_: esiRegistered', 9, ('PayRoll Rates', 'E', 'ESI Enrolled by Default'))

# --- this month's inputs ---------------------------------------------------------------------
col('miRow', 'Monthly Inputs row', W('IFERROR(MATCH(MK&"|"&{code},%s,0),"")' % rng('Monthly Inputs', 'Q')), 'input', '0',
    'Row on Monthly Inputs holding this month\'s entry for this person (blank = nothing entered).', 'Monthly Inputs › Month + Employee Code (helper col Q)', 'payForMonth_(ss,\'Monthly Inputs\',month)', 9)
col('inPaid', 'Paid Days typed in', W('IF({miRow}="","",IF(%s="","",%s))' % (MI('C'), num(MI('C')))), 'input', '0.00', 'Paid days entered (leave upload writes it). Blank = full month, or days employed for a joiner/leaver.', 'Monthly Inputs › Paid Days (col C)', "inp['Paid Days']", 9, ('Monthly Inputs', 'C', 'Paid Days'))
col('clDays', 'CL days encashed', W('IF({miRow}="",0,%s)' % num(MI('D'))), 'input', '0.00', 'CL days paid out in cash this month.', 'Monthly Inputs › CL Days Encashed (col D)', "Number(input.clDays)||0", 9, ('Monthly Inputs', 'D', 'CL Days Encashed'))
col('holdTxt', 'Hold (F&F/Grievance)', W('IF({miRow}="","",TRIM(%s))' % MI('L')), 'input', None, 'Hold type typed on the page.', 'Monthly Inputs › Hold (F&F/Grievance) (col L)', "inp['Hold (F&F/Grievance)']", 11, ('Monthly Inputs', 'L', 'Hold (F&F/Grievance)'))
col('holdOld', 'Hold for F&F (old Y/N)', W('IF({miRow}="","",TRIM(%s))' % MI('E')), 'input', None, 'Older Y/N hold column, still honoured.', 'Monthly Inputs › Hold for F&F (Y/N) (col E)', "payYes_(inp['Hold for F&F (Y/N)'])", 9, ('Monthly Inputs', 'E', 'Hold for F&F (Y/N)'))
col('release', 'Release held (₹)', W('IF({miRow}="",0,%s)' % num(MI('M'))), 'input', '#,##0', 'Earlier held salary paid out this month (added to bank payable only).', 'Monthly Inputs › Release Held (₹) (col M)', "Number(input.release)||0", 10, ('Monthly Inputs', 'M', 'Release Held (₹)'))
col('skipLoan', 'Skip loan recovery (1/0)', W('IF({miRow}="",0,%s)' % yes(MI('N'))), 'input', '0', '1 = this month\'s advance recovery is skipped (reason on the page).', 'Monthly Inputs › Skip Loan Recovery (Y/N) (col N)', "payYes_(inp['Skip Loan Recovery (Y/N)'])", 9, ('Monthly Inputs', 'N', 'Skip Loan Recovery (Y/N)'))
col('holdType', 'Hold type', W('IF(LEFT(LOWER({holdTxt}),1)="f","fnf",IF(LEFT(LOWER({holdTxt}),1)="g","grievance",IF(LEFT(LOWER({holdOld}),1)="y","fnf","")))'), 'calc', None,
    'fnf = salary held for Full & Final; grievance = withheld; blank = paid normally.', 'derived', 'payCompute_: hold', 9)

# --- paid days and earnings ------------------------------------------------------------------
col('employed', 'Days employed this month', W('MAX(0,IF(AND({lwd}<>"",{lwd}<ME),{lwd},ME)-IF({doj}>MS,{doj},MS)+1)'), 'calc', '0',
    'From the later of joining date / 1st to the earlier of last working day / month end, both days counted.', 'Date of Joining, Last Working Day', 'payEmployedDays_', 10)
col('paid', 'Paid days', W('IF({inPaid}="",IF({employed}<DIM,{employed},DIM),{inPaid})'), 'calc', '0.00',
    'Paid days typed in; if blank, the full month (or days employed for a joiner/leaver).', 'Paid Days typed in / Days employed / days in month', "payCompute_: paidDays", 8)
col('basic', 'Month Basic', W(rnd('{fullBasic}*{paid}/DIM')), 'calc', '#,##0', 'ROUND(Full Basic × Paid Days ÷ Days in month).', '', "payCalc_: basic = R(fullBasic*paidDays/daysInMonth)", 10)
col('adaAmt', 'ADA', W(rnd('{basic}*{adaPct}')), 'calc', '#,##0', 'ROUND(Month Basic × ADA %).', '', 'payCalc_: ada', 9)
col('daAmt', 'DA', W(rnd('{basic}*{daPct}')), 'calc', '#,##0', 'ROUND(Month Basic × DA %).', '', 'payCalc_: da', 9)
col('wages', 'Wages (Basic+ADA+DA)', W('{basic}+{adaAmt}+{daAmt}'), 'calc', '#,##0', 'EPF / ESI wage base. Tuition and CL encashment are NOT part of it.', '', 'payCalc_: wages', 11)

# --- EPF / ESI ---------------------------------------------------------------------------------
col('epf', 'EPF (employee = employer)', W('IF(AND({epfMem}=1,{epfReg}=1),%s,0)' % up('MIN({wages},EPF_CEIL)*EPF_RATE')), 'calc', '#,##0',
    '12% of Wages capped at the EPF ceiling, rounded UP. Only EPF members at EPF-registered schools. Employer pays the same amount.', 'Settings: EPF Rate % + PF Rate %, EPF Cap Salary', 'payCalc_: epf = up(min(wages, epfCeiling) × epfRate)', 11)
col('esiEffMax', 'ESI: latest Salary Master date on/before period start', W('_xlfn.MAXIFS(%s,%s,{code},%s,"<="&PS)' % (rng('Salary Master', 'V'), rng('Salary Master', 'Y'), rng('Salary Master', 'V'))), 'calc', 'dd-mmm-yyyy',
    'ESI coverage is fixed for the 6-month period (Apr-Sep / Oct-Mar) by the pay in force at its start. This is the Effective From of that row (0 = joined after the period began).', 'Salary Master › Effective From (helper col V)', 'payEsiCovered_', 12)
col('esiEff', 'ESI: Salary Master date used', W('IF({esiEffMax}=0,_xlfn.MINIFS(%s,%s,{code}),{esiEffMax})' % (rng('Salary Master', 'V'), rng('Salary Master', 'Y'))), 'calc', 'dd-mmm-yyyy',
    'If joined after the period began, the first row they ever had.', 'Salary Master › Effective From', 'payEsiCovered_', 12)
col('esiRow', 'ESI: Salary Master row used', W('IFERROR(MATCH({code}&"|"&TEXT({esiEff},"yyyymmdd"),%s,0),"")' % rng('Salary Master', 'U')), 'calc', '0', 'Sheet row on Salary Master used to decide ESI coverage.', 'Salary Master (helper col U)', 'payEsiCovered_', 9)
col('esiRateRow', 'ESI: Rates row used', W('IF({esiRow}="","",IFERROR(MATCH(SUBSTITUTE(UPPER(INDEX(%s,{esiRow}))," ",""),%s,0),""))' % (rng('Salary Master', 'C'), rng('PayRoll Rates', 'F'))), 'calc', '0', "PayRoll Rates row for that old row's school.", 'PayRoll Rates', 'payEsiCovered_', 9)
col('esiBasic', 'ESI: Full Basic then', W('IF({esiRow}="","",%s)' % num("INDEX(%s,{esiRow})" % rng('Salary Master', 'G'))), 'calc', '#,##0', 'Full Basic on that old row.', 'Salary Master › Full Basic (col G) at the ESI row', 'payEsiCovered_', 10)
col('esiWagesFull', 'ESI: full-month wages then', W('IF({esiRateRow}="","",{esiBasic}+ROUND(ROUND({esiBasic}*%s/100,6),0)+ROUND(ROUND({esiBasic}*%s/100,6),0))' %
    (num("INDEX(%s,{esiRateRow})" % rng('PayRoll Rates', 'C')), num("INDEX(%s,{esiRateRow})" % rng('PayRoll Rates', 'B')))), 'calc', '#,##0',
    'Basic + ADA + DA for a FULL month on that row (not prorated).', 'PayRoll Rates › ADA % / DA %', 'payEsiCovered_', 11)
col('esiCovered', 'ESI covered this period?', W('IF({esiWagesFull}="",{wages}<=ESI_THR,{esiWagesFull}<=ESI_THR)'), 'calc', None,
    'TRUE if those full-month wages are at or under the ESI threshold (₹21,000). If the old row cannot be found, this month\'s wages are used instead.', 'Settings: ESI Threshold', 'payEsiCovered_', 10)
col('esi', 'ESI employee', W('IF(AND({esiReg}=1,{esiCovered}),%s,0)' % up('{wages}*ESI_EE')), 'calc', '#,##0', 'Wages × 0.75%, rounded UP; only if the school is ESI-registered and the person is covered.', 'Settings: ESI Employee %', 'payCalc_: esi', 10)
col('esiEr', 'ESI employer', W('IF(AND({esiReg}=1,{esiCovered}),%s,0)' % up('{wages}*ESI_ER')), 'calc', '#,##0', 'Wages × 3.25%, rounded UP; employer cost, not deducted.', 'Settings: ESI Employer %', 'payCalc_: esiEr', 10)

# --- pay lines ---------------------------------------------------------------------------------
col('clEnc', 'CL encashment', W(rnd('{clDays}*({wages}-{epf}-{esi})/CL_DIV')), 'calc', '#,##0', 'CL days × (Wages − EPF − ESI) ÷ divisor (30).', 'Settings: CL Encashment Divisor', 'payCalc_: clEnc', 10)
AD = lambda cond: 'SUMIFS(%s,%s,MK&"|"&{code}%s)' % (rng('Adjustments', 'F'), rng('Adjustments', 'K'), cond)
col('adjE', 'Other earnings (Adjustments)', W(AD(',%s,"Earning",%s,"<>TDS"' % (rng('Adjustments', 'E'), rng('Adjustments', 'D')))), 'input', '#,##0',
    'Sum of this month\'s Earning adjustments (arrears, allowances, bonus...).', 'Adjustments › Amount (col F) where Month+Code match and Direction = Earning', "payCompute_: sumDir('Earning')", 11, ('Adjustments', 'F', 'Amount (Earning rows)'))
col('adjD', 'Other deductions (Adjustments)', W(AD(',%s,"Deduction",%s,"<>TDS"' % (rng('Adjustments', 'E'), rng('Adjustments', 'D')))), 'input', '#,##0',
    'Sum of this month\'s Deduction adjustments except TDS (fines, fee recovery, manual advance recovery...).', 'Adjustments › Amount (col F) where Direction = Deduction and Type is not TDS', "payCompute_: sumDir('Deduction')", 11, ('Adjustments', 'F', 'Amount (Deduction rows)'))
col('adjTds', 'TDS (Adjustments)', W(AD(',%s,"TDS"' % rng('Adjustments', 'D'))), 'input', '#,##0', 'TDS entered as an adjustment this month.', 'Adjustments › Amount where Type = TDS', 'payCompute_: tds', 10, ('Adjustments', 'F', 'Amount (TDS rows)'))
col('tds', 'TDS total', W('{tdsMaster}+{adjTds}'), 'calc', '#,##0', 'Standing monthly TDS + TDS adjustments.', '', 'payCompute_: tds', 9)
col('gross', 'Gross', W('{wages}+{clEnc}+{tuition}+{adjE}'), 'result', '#,##0', 'Wages + CL encashment + tuition + other earnings.', '', 'payCalc_: gross', 11)
col('cti', 'CTI', W('{gross}+{epf}+{esiEr}'), 'result', '#,##0', 'Cost to institution = Gross + employer EPF + employer ESI.', '', 'payCalc_: cti', 11)

# --- RRF ---------------------------------------------------------------------------------------
LG = lambda acct: 'SUMIFS(%s,%s,{code},%s,"%s",%s,"<"&MS)' % (rng('Ledger', 'F'), rng('Ledger', 'C'), rng('Ledger', 'D'), acct, rng('Ledger', 'K'))
col('months', 'Months of service', W('IF({doj}>ME,0,DATEDIF({doj},ME,"m"))'), 'calc', '0', 'Complete months from joining to month end.', 'Date of Joining', 'payMonthsBetween_', 9)
col('openRrf', 'RRF balance at start of month', W(LG('RRF')), 'input', '#,##0', 'Sum of RRF ledger entries dated before this month.', 'Ledger › Amount (col F) where Account = RRF and Month is before this month', 'payLedgerBalances_(ss,\'RRF\',month)', 11, ('Ledger', 'F', 'Amount (RRF rows)'))
col('rrfTarget', 'RRF target', W('RRF_TGT*{cti}'), 'calc', '#,##0', 'Target months (3) × CTI. Only matters when "RRF Stop At Target" = 1.', 'Settings: RRF Target Months', 'payCalc_: target', 10)
col('rrfRate', 'RRF rate', W('IF({rrfMem}=1,IF({months}<12,RRF1,IF({months}<24,RRF2,IF(OR({months}<36,AND(RRF_STOP=1,{openRrf}<{rrfTarget})),RRF3,0))),0)'), 'calc', '0%',
    '12% in year 1, 9% year 2, 6% year 3; after 36 months 0%, unless "Stop At Target" is on and the balance is below target (then 6%).', 'Settings: RRF Y1/Y2/Y3 %', 'payCalc_: rrfRate', 8)
col('rrf', 'RRF deduction', W('IF(AND({rrfRate}>0,RRF_STOP=1),MIN(%s,MAX(0,%s)),%s)' % (rnd('{rrfRate}*{cti}'), rnd('{rrfTarget}-{openRrf}'), rnd('{rrfRate}*{cti}'))), 'calc', '#,##0',
    'Rate × CTI; with "Stop At Target" on it is capped at the gap to target.', '', 'payCalc_: rrf', 10)

# --- loans -------------------------------------------------------------------------------------
col('loanSched', 'Loan: scheduled recovery', W('SUMIFS(%s,%s,{code},%s,"<>closed*",%s,"<="&MS)' % (rng('Loans', 'F'), rng('Loans', 'B'), rng('Loans', 'H'), rng('Loans', 'M'))), 'input', '#,##0',
    'Monthly Recovery of every loan not Closed whose Start Month has begun.', 'Loans › Monthly Recovery (col F) where Employee Code matches, Status is not Closed, Start Month <= this month', 'payLoanSchedules_', 11, ('Loans', 'F', 'Monthly Recovery'))
col('loanBal', 'Loan: balance still owed', W(LG('Loan')), 'input', '#,##0', 'Sum of Loan ledger entries dated before this month.', 'Ledger › Amount (col F) where Account = Loan and Month is before this month', "payLedgerBalances_(ss,'Loan',month)", 11, ('Ledger', 'F', 'Amount (Loan rows)'))
col('loanDue', 'Loan: due', W('MAX(0,MIN({loanSched},{loanBal}))'), 'calc', '#,##0', 'Scheduled recovery, never more than the balance.', '', 'payCompute_: loanDue', 9)
col('loanRec', 'Loan recovery', W('MAX(0,MIN(IF({skipLoan}=1,0,{loanDue}),{gross}-({epf}+{esi}+{rrf}+{tuition}+{adjD}+{tds})))'), 'calc', '#,##0',
    'Loan due (0 if skipped), never more than what is left of gross after the other deductions.', '', 'payCalc_: loanRecovery', 10)

# --- results -----------------------------------------------------------------------------------
col('totDed', 'Total deductions', W('{epf}+{esi}+{rrf}+{tuition}+{adjD}+{tds}+{loanRec}'), 'result', '#,##0', 'EPF + ESI + RRF + tuition + other deductions + TDS + loan recovery.', '', 'payCalc_: totalDeductions', 11)
col('net', 'NET PAY', W('{gross}-{totDed}'), 'result', '#,##0', 'Gross − total deductions.', '', 'payCalc_: net', 11)
col('bank', 'Bank payable', W('IF({holdType}<>"",0,{net})+{release}'), 'result', '#,##0', 'Net pay (0 if the salary is held) + any earlier held salary released this month.', '', 'payCalc_: bankPayable', 11)
col('heldF', 'Held for F&F', W('IF({holdType}="fnf",{net},0)'), 'result', '#,##0', 'Net pay kept back for Full & Final.', '', 'payCalc_: heldForFnF', 10)
col('withheld', 'Withheld (grievance)', W('IF({holdType}="grievance",{net},0)'), 'result', '#,##0', 'Net pay withheld for a grievance.', '', 'payCalc_: withheld', 10)
col('grat', 'Gratuity provision', W(rnd('({basic}+{daAmt})*GRAT')), 'result', '#,##0', '5% of (Basic + DA) -- an employer provision, not deducted.', 'Settings: Gratuity Provision %', 'payCalc_: gratuityProvision', 10)

# --- cross-check vs the portal's locked register -----------------------------------------------
RG = lambda c: "INDEX(%s,{regRow})" % rng('Payroll Register', c)
col('regRow', 'Register row', 'IFERROR(MATCH(MK&"|"&{code},%s,0),"")' % rng('Payroll Register', 'AI'), 'check', '0', 'Row on the Payroll Register tab (locked months only).', 'Payroll Register › Month + Employee Code', 'payLock_ writes the register', 9, ('Payroll Register', 'U', 'Net Pay'))
col('regNet', 'Register Net Pay', 'IF({regRow}="","",%s)' % num(RG('U')), 'check', '#,##0', 'What the portal locked.', 'Payroll Register › Net Pay (col U)', '', 11, ('Payroll Register', 'U', 'Net Pay'))
col('diffNet', 'Net: sheet − register', 'IF(OR({regRow}="",{incl}<>"Yes"),"",{net}-{regNet})', 'check', '#,##0;[Red]-#,##0;0', 'Should be 0 for every locked employee.', '', '', 11)
col('regGross', 'Register Gross', 'IF({regRow}="","",%s)' % num(RG('M')), 'check', '#,##0', '', 'Payroll Register › Gross (col M)', '', 11, ('Payroll Register', 'M', 'Gross'))
col('diffGross', 'Gross: sheet − register', 'IF(OR({regRow}="",{incl}<>"Yes"),"",{gross}-{regGross})', 'check', '#,##0;[Red]-#,##0;0', 'Should be 0.', '', '', 11)
col('regCti', 'Register CTI', 'IF({regRow}="","",%s)' % num(RG('Z')), 'check', '#,##0', '', 'Payroll Register › CTI (col Z)', '', 11, ('Payroll Register', 'Z', 'CTI'))
col('diffCti', 'CTI: sheet − register', 'IF(OR({regRow}="",{incl}<>"Yes"),"",{cti}-{regCti})', 'check', '#,##0;[Red]-#,##0;0', 'Should be 0.', '', '', 11)
col('ok', 'Matches register?', 'IF({regRow}="","not locked",IF({incl}<>"Yes","n/a",IF(AND({diffNet}=0,{diffGross}=0,{diffCti}=0),"OK","DIFFERENT")))', 'check', None, 'OK when net, gross and CTI all agree with the locked register.', '', '', 12)

KEYS = [c['key'] for c in COLS]
LET = {k: L(i + 1) for i, k in enumerate(KEYS)}
HDR_ROW = 24
FIRST = HDR_ROW + 1


def fill(formula, r):
    if formula is None:
        return None
    f = formula if formula.startswith('=') else '=' + formula
    return re.sub(r'\{(\w+)\}', lambda m: '%s%d' % (LET[m.group(1)], r), f)


# ---------------------------------------------------------------------------------------------
def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and re.match(r'\d{4}-\d{2}-\d{2}', v.strip()):
        return dt.date.fromisoformat(v.strip()[:10])
    return None


def month_key(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime('%Y-%m')
    return str(v or '').strip()[:7]


def build(path, data, month, banner):
    """data: {tab: [[cells], ...]} WITHOUT the header row (headers come from HEADERS).
    month: datetime.date (any day -- first of that month is used)."""
    wb = Workbook()
    month = dt.date(month.year, month.month, 1)
    ws_help = wb.active
    ws_help.title = 'Read Me'
    calc = wb.create_sheet('Salary Calc')
    summ = wb.create_sheet('School Summary')
    guide = wb.create_sheet('Formula Guide')
    inputs = {}
    for tab in ['PayRoll Constants', 'PayRoll Rates', 'Salary Master', 'Monthly Inputs', 'Adjustments', 'Ledger', 'Loans', 'Payroll Register']:
        inputs[tab] = wb.create_sheet(tab)
    sm_rows = data['Salary Master']

    # ---- input tabs -------------------------------------------------------------------------
    for tab, ws in inputs.items():
        hdr = HEADERS[tab]
        for j, h in enumerate(hdr, 1):
            c = ws.cell(1, j, h)
            c.font = HDR_FONT
            c.fill = HDR_FILL
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.freeze_panes = 'A2'
        for i, row in enumerate(data.get(tab, []), 2):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                c.font = BLUE
                if isinstance(v, (dt.date, dt.datetime)):
                    c.number_format = 'yyyy-mm-dd'
        for j in range(1, len(hdr) + 1):
            ws.column_dimensions[L(j)].width = 16
    # helper columns
    def helpers(tab, rows):
        ws = inputs[tab]
        for r in range(2, LIMIT[tab] + 1):
            for col_l, f in rows.items():
                c = ws[f'{col_l}{r}']
                c.value = '=' + f.replace('#', str(r))
                c.font = Font(color='7F7F7F', italic=True)
    sm = 'Salary Master'
    helpers(sm, {
        'U': 'IF(OR(Y#="",V#=""),"",Y#&"|"&TEXT(V#,"yyyymmdd"))',
        'V': 'IF(F#="","",IF(ISNUMBER(F#),F#,IFERROR(DATEVALUE(LEFT(TRIM(F#),10)),"")))',
        'W': 'IF(E#="","",IF(ISNUMBER(E#),E#,IFERROR(DATEVALUE(LEFT(TRIM(E#),10)),"")))',
        'X': 'IF(R#="","",IF(ISNUMBER(R#),R#,IFERROR(DATEVALUE(LEFT(TRIM(R#),10)),"")))',
        'Y': 'TRIM(A#)',
    })
    for cl, h in {'U': 'Key (code|Effective From)', 'V': 'Effective From (date)', 'W': 'Date of Joining (date)', 'X': 'Last Working Day (date)', 'Y': 'Code (trimmed)'}.items():
        c = inputs[sm][f'{cl}1']
        c.value = h
        c.font = Font(bold=True, color='7F7F7F')
    helpers('Monthly Inputs', {'Q': 'IF(B#="","",IF(ISNUMBER(A#),TEXT(A#,"yyyy-mm"),LEFT(TRIM(A#),7))&"|"&TRIM(B#))'})
    inputs['Monthly Inputs']['Q1'].value = 'Key (month|code)'
    helpers('Adjustments', {'K': 'IF(C#="","",IF(ISNUMBER(B#),TEXT(B#,"yyyy-mm"),LEFT(TRIM(B#),7))&"|"&TRIM(C#))'})
    inputs['Adjustments']['K1'].value = 'Key (month|code)'
    helpers('Ledger', {'K': 'IF(B#="","",IF(ISNUMBER(B#),DATE(YEAR(B#),MONTH(B#),1),IFERROR(DATEVALUE(LEFT(TRIM(B#),7)&"-01"),"")))'})
    inputs['Ledger']['K1'].value = 'Month (1st of month, date)'
    helpers('Loans', {'M': 'IF(B#="","",IF(G#="",0,IF(ISNUMBER(G#),DATE(YEAR(G#),MONTH(G#),1),IFERROR(DATEVALUE(LEFT(TRIM(G#),7)&"-01"),0))))'})
    inputs['Loans']['M1'].value = 'Start Month (1st of month, date)'
    helpers('Payroll Register', {'AI': 'IF(B#="","",IF(ISNUMBER(A#),TEXT(A#,"yyyy-mm"),LEFT(TRIM(A#),7))&"|"&TRIM(B#))'})
    inputs['Payroll Register']['AI1'].value = 'Key (month|code)'
    helpers('PayRoll Rates', {'F': 'IF(A#="","",SUBSTITUTE(UPPER(A#)," ",""))'})
    inputs['PayRoll Rates']['F1'].value = 'School key'
    for tab, cl in [('Monthly Inputs', 'Q'), ('Adjustments', 'K'), ('Ledger', 'K'), ('Loans', 'M'), ('Payroll Register', 'AI'), ('PayRoll Rates', 'F')]:
        inputs[tab][f'{cl}1'].font = Font(bold=True, color='7F7F7F')
    for tab in ('Salary Master', 'Monthly Inputs', 'Adjustments', 'Ledger', 'Loans', 'Payroll Register', 'PayRoll Rates'):
        inputs[tab].sheet_properties.tabColor = 'FFC000'
    inputs['PayRoll Constants'].sheet_properties.tabColor = 'FFC000'

    # ---- Salary Calc: parameters --------------------------------------------------------------
    c = calc
    c['A1'] = 'Salary Calc -- one row per employee, every formula live'
    c['A1'].font = Font(bold=True, size=14)
    c['A2'] = banner
    c['A2'].font = Font(bold=True, color='C00000')
    c['A4'], c['B4'] = 'Month (type the 1st of the month)', month
    c['B4'].number_format = 'dd-mmm-yyyy'
    c['B4'].font = Font(bold=True, color='0000CC')
    c['B4'].fill = PatternFill('solid', fgColor='FFFF00')
    c['A5'], c['B5'] = 'Month end', '=EOMONTH(B4,0)'
    c['A6'], c['B6'] = 'Days in month', '=DAY(B5)'
    c['A7'], c['B7'] = 'Month key', '=TEXT(B4,"yyyy-mm")'
    c['A8'], c['B8'] = 'ESI contribution period starts', '=IF(AND(MONTH(B4)>=4,MONTH(B4)<=9),DATE(YEAR(B4),4,1),DATE(IF(MONTH(B4)>=10,YEAR(B4),YEAR(B4)-1),10,1))'
    c['B5'].number_format = c['B8'].number_format = 'dd-mmm-yyyy'
    for nm, ref in [('MS', '$B$4'), ('ME', '$B$5'), ('DIM', '$B$6'), ('MK', '$B$7'), ('PS', '$B$8')]:
        wb.defined_names[nm] = DefinedName(nm, attr_text="'Salary Calc'!" + ref)
    def pc(name, default=None):
        look = num('INDEX(%s,MATCH("%s",%s,0))' % (rng('PayRoll Constants', 'B'), name, rng('PayRoll Constants', 'A')))
        if default is None:
            return look
        # the backend appends these four settings with a default when the row is missing
        return 'IF(ISNUMBER(MATCH("%s",%s,0)),%s,%s)' % (name, rng('PayRoll Constants', 'A'), look, default)
    settings = [
        ('EPF_RATE', 'EPF rate (employee = employer)', '=(%s+%s)/100' % (pc('EPF Rate %'), pc('PF Rate %')), '0.00%', 'PayRoll Constants: "EPF Rate %" (8.33, the EPS share) + "PF Rate %" (3.67) = 12%', ['EPF Rate %', 'PF Rate %']),
        ('EPF_CEIL', 'EPF wage ceiling', '=' + pc('EPF Cap Salary'), '#,##0', 'PayRoll Constants: "EPF Cap Salary"', ['EPF Cap Salary']),
        ('ESI_EE', 'ESI employee %', '=%s/100' % pc('ESI Employee %'), '0.00%', 'PayRoll Constants: "ESI Employee %"', ['ESI Employee %']),
        ('ESI_ER', 'ESI employer %', '=%s/100' % pc('ESI Employer %'), '0.00%', 'PayRoll Constants: "ESI Employer %"', ['ESI Employer %']),
        ('ESI_THR', 'ESI wage threshold', '=' + pc('ESI Threshold'), '#,##0', 'PayRoll Constants: "ESI Threshold"', ['ESI Threshold']),
        ('RRF1', 'RRF year 1', '=%s/100' % pc('RRF Y1 %'), '0%', 'PayRoll Constants: "RRF Y1 %"', ['RRF Y1 %']),
        ('RRF2', 'RRF year 2', '=%s/100' % pc('RRF Y2 %'), '0%', 'PayRoll Constants: "RRF Y2 %"', ['RRF Y2 %']),
        ('RRF3', 'RRF year 3 (and after, if below target)', '=%s/100' % pc('RRF Y3 %'), '0%', 'PayRoll Constants: "RRF Y3 %"', ['RRF Y3 %']),
        ('RRF_TGT', 'RRF target (months of CTI)', '=' + pc('RRF Target Months', 3), '0', 'PayRoll Constants: "RRF Target Months" (3 if the row is absent -- the backend adds it)', ['RRF Target Months']),
        ('RRF_STOP', 'RRF Stop At Target (1 = on)', '=IF(%s=1,1,0)' % pc('RRF Stop At Target (1=Yes)', 0), '0', 'PayRoll Constants: "RRF Stop At Target (1=Yes)" (only exactly 1 counts)', ['RRF Stop At Target (1=Yes)']),
        ('GRAT', 'Gratuity provision %', '=%s/100' % pc('Gratuity Provision %', 5), '0.0%', 'PayRoll Constants: "Gratuity Provision %"', ['Gratuity Provision %']),
        ('CL_DIV', 'CL encashment divisor', '=' + pc('CL Encashment Divisor', 30), '0', 'PayRoll Constants: "CL Encashment Divisor"', ['CL Encashment Divisor']),
    ]
    c['D3'], c['E3'], c['F3'] = 'Setting (from PayRoll Constants)', 'Value', 'Constants row'
    for cc in ('D3', 'E3', 'F3'):
        c[cc].font = BOLD
    pc_rows = {str(r[0]).strip(): i for i, r in enumerate(data['PayRoll Constants'], 2) if r and r[0]}
    for i, (nm, label, f, fmt, srcnote, names) in enumerate(settings, 4):
        c[f'D{i}'] = label
        c[f'E{i}'] = f
        c[f'E{i}'].number_format = fmt
        c[f'F{i}'] = '=IFERROR(MATCH("%s",%s,0),"not found")' % (names[0], rng('PayRoll Constants', 'A'))
        found = ', '.join('"%s" = row %d' % (n, pc_rows[n]) if n in pc_rows else '"%s" not on the tab' % n for n in names)
        c[f'E{i}'].comment = Comment('Fetched from the PayRoll Constants tab (LMCS-Salary-PayScale workbook).\n%s\nAt build time: %s' % (srcnote, found), 'Payroll proofreading')
        wb.defined_names[nm] = DefinedName(nm, attr_text="'Salary Calc'!$E$%d" % i)

    # ---- Salary Calc: table ---------------------------------------------------------------------
    for j, cd in enumerate(COLS, 1):
        h = calc.cell(HDR_ROW, j, cd['header'])
        h.font = HDR_FONT
        h.fill = PatternFill('solid', fgColor={'who': '1F3864', 'master': '375623', 'input': '7F6000', 'calc': '3A3A3A', 'result': '833C0B', 'check': '595959'}[cd['group']])
        h.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        calc.column_dimensions[L(j)].width = cd['width']
        tip = cd['desc']
        if cd['src']:
            tip += '\n\nSOURCE: ' + cd['src']
        if cd['code']:
            tip += '\nBACKEND: ' + cd['code']
        h.comment = Comment(tip, 'Payroll proofreading', width=360, height=170)
    calc.row_dimensions[HDR_ROW].height = 62
    calc.freeze_panes = calc.cell(FIRST, 4)

    # python mirror of the row lookups, only to write exact row numbers into the cell comments
    ms_by_code = {}
    for i, r in enumerate(sm_rows, 2):
        code = str(r[0] or '').strip()
        eff = to_date(r[5])
        if code and eff:
            ms_by_code.setdefault(code, []).append((eff, i, r))
    me = (month.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
    mk = month.strftime('%Y-%m')

    def rows_for(tab, pred):
        out = []
        for i, r in enumerate(data.get(tab, []), 2):
            try:
                if pred(r):
                    out.append(i)
            except Exception:
                pass
        return out
    mi_rows = {}
    for i, r in enumerate(data.get('Monthly Inputs', []), 2):
        if r[1]:
            mi_rows[month_key(r[0]) + '|' + str(r[1]).strip()] = i
    rates_rows = {re.sub(r'\s+', '', str(r[0] or '').upper()): i for i, r in enumerate(data['PayRoll Rates'], 2) if r and r[0]}
    codes = []
    for r in sm_rows:
        k = str(r[0] or '').strip()
        if k and k not in codes:
            codes.append(k)
    n = len(codes)
    for idx, code in enumerate(codes):
        r = FIRST + idx
        for j, cd in enumerate(COLS, 1):
            cell = calc.cell(r, j)
            if cd['key'] == 'code':
                cell.value = code
                cell.font = BLUE
            else:
                cell.value = fill(cd['formula'], r)
            if cd['fmt']:
                cell.number_format = cd['fmt']
            if cd['group'] in ('result',):
                cell.fill = PatternFill('solid', fgColor=GROUPS['result'])
            elif cd['group'] == 'input':
                cell.fill = PatternFill('solid', fgColor=GROUPS['input'])
            elif cd['group'] == 'master':
                cell.fill = PatternFill('solid', fgColor=GROUPS['master'])
            elif cd['group'] == 'who':
                cell.fill = PatternFill('solid', fgColor=GROUPS['who'])
            if cd['key'] == 'net':
                cell.font = BOLD
        # exact-row source comments (as of the month loaded when the file was built)
        cands = [x for x in ms_by_code.get(code, []) if x[0] <= me]
        pick = max(cands, key=lambda x: (x[0], x[1])) if cands else None
        def note(key, text):
            calc.cell(r, KEYS.index(key) + 1).comment = Comment(text + '\n\n(row numbers are for the month the file was built with: %s; the "row used" columns recalculate live)' % mk, 'Payroll proofreading', width=330, height=120)
        if pick:
            srow = pick[1]
            hdr = HEADERS['Salary Master']
            for key, colix in [('name', 1), ('entRaw', 2), ('desig', 3), ('doj', 4), ('lwd', 17), ('status', 15), ('fullBasic', 6), ('epfMem', 7), ('rrfMem', 8), ('tuition', 9), ('tdsMaster', 18)]:
                note(key, 'Salary Master tab, row %d, column %s "%s"\n(this code has %d row(s) on the tab; row %d is the latest with Effective From %s on/before %s)' %
                     (srow, L(colix + 1), hdr[colix], len(ms_by_code[code]), srow, pick[0], me))
            ent = re.sub(r'\s+', '', str(pick[2][2] or '').upper())
            rr = rates_rows.get(ent)
            for key, cl, h in [('daPct', 'B', 'DA %'), ('adaPct', 'C', 'ADA %'), ('epfReg', 'D', 'EPF Enrolled by Default'), ('esiReg', 'E', 'ESI Enrolled by Default')]:
                note(key, 'PayRoll Rates tab, row %s, column %s "%s" (school "%s")' % (rr or 'NOT FOUND', cl, h, pick[2][2]))
            mrow = mi_rows.get(mk + '|' + code)
            for key, cl, h in [('inPaid', 'C', 'Paid Days'), ('clDays', 'D', 'CL Days Encashed'), ('holdTxt', 'L', 'Hold (F&F/Grievance)'), ('holdOld', 'E', 'Hold for F&F (Y/N)'), ('release', 'M', 'Release Held (₹)'), ('skipLoan', 'N', 'Skip Loan Recovery (Y/N)')]:
                note(key, ('Monthly Inputs tab, row %d, column %s "%s"' % (mrow, cl, h)) if mrow else 'Monthly Inputs tab: no row for %s / %s -> treated as blank/0' % (mk, code))
            ar = rows_for('Adjustments', lambda x: month_key(x[1]) == mk and str(x[2]).strip() == code)
            for key, what in [('adjE', 'Earning'), ('adjD', 'Deduction'), ('adjTds', 'TDS')]:
                sel = [i for i in ar if (str(data['Adjustments'][i - 2][4]) == what and str(data['Adjustments'][i - 2][3]) != 'TDS') or (what == 'TDS' and str(data['Adjustments'][i - 2][3]) == 'TDS')] if ar else []
                note(key, ('Adjustments tab, row(s) %s, column F "Amount" (%s)' % (', '.join(map(str, sel)), what)) if sel else 'Adjustments tab: no %s rows for %s / %s' % (what, mk, code))
            lr = rows_for('Ledger', lambda x: str(x[2]).strip() == code and str(x[3]).strip().lower() == 'rrf' and month_key(x[1]) < mk)
            note('openRrf', ('Ledger tab, row(s) %s, column F "Amount" where Account = RRF and Month is before %s' % (', '.join(map(str, lr)), mk)) if lr else 'Ledger tab: no RRF entries before %s for %s' % (mk, code))
            lr = rows_for('Ledger', lambda x: str(x[2]).strip() == code and str(x[3]).strip().lower() == 'loan' and month_key(x[1]) < mk)
            note('loanBal', ('Ledger tab, row(s) %s, column F "Amount" where Account = Loan and Month is before %s' % (', '.join(map(str, lr)), mk)) if lr else 'Ledger tab: no Loan entries before %s for %s' % (mk, code))
            ln = rows_for('Loans', lambda x: str(x[1]).strip() == code and not str(x[7] or '').lower().startswith('closed'))
            note('loanSched', ('Loans tab, row(s) %s, column F "Monthly Recovery" (not Closed; counted only once Start Month has begun)' % ', '.join(map(str, ln))) if ln else 'Loans tab: no open loans for %s' % code)
            rg = rows_for('Payroll Register', lambda x: month_key(x[0]) == mk and str(x[1]).strip() == code)
            for key, cl, h in [('regNet', 'U', 'Net Pay'), ('regGross', 'M', 'Gross'), ('regCti', 'Z', 'CTI')]:
                note(key, ('Payroll Register tab, row %d, column %s "%s"' % (rg[0], cl, h)) if rg else 'Payroll Register tab: %s is not locked for %s (no row)' % (mk, code))

    last = FIRST + n - 1
    calc.auto_filter.ref = '%s%d:%s%d' % ('A', HDR_ROW, L(len(COLS)), last)
    calc.conditional_formatting.add('%s%d:%s%d' % (LET['ok'], FIRST, LET['ok'], last), CellIsRule(operator='equal', formula=['"DIFFERENT"'], fill=PatternFill('solid', bgColor='FF9999')))
    calc.conditional_formatting.add('%s%d:%s%d' % (LET['ok'], FIRST, LET['ok'], last), CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', bgColor='C6EFCE')))
    calc.conditional_formatting.add('%s%d:%s%d' % (LET['incl'], FIRST, LET['incl'], last), CellIsRule(operator='equal', formula=['"No"'], fill=PatternFill('solid', bgColor='D9D9D9')))
    # totals row above the table
    tot = HDR_ROW - 1
    calc.cell(tot, 1, 'TOTAL (included staff)').font = BOLD
    for key in ['fullBasic', 'basic', 'adaAmt', 'daAmt', 'wages', 'epf', 'esi', 'esiEr', 'clEnc', 'adjE', 'adjD', 'tds', 'gross', 'cti', 'rrf', 'loanRec', 'totDed', 'net', 'bank', 'heldF', 'withheld', 'grat']:
        cl = LET[key]
        cc = calc[f'{cl}{tot}']
        cc.value = '=SUM(%s%d:%s%d)' % (cl, FIRST, cl, last)
        cc.number_format = '#,##0'
        cc.font = BOLD
    calc.cell(tot - 1, 1, 'Headcount on payroll').font = BOLD
    calc.cell(tot - 1, 2, '=COUNTIF(%s%d:%s%d,"Yes")' % (LET['incl'], FIRST, LET['incl'], last)).font = BOLD

    # ---- School Summary -------------------------------------------------------------------------
    summ['A1'] = 'School-wise totals for the month on the Salary Calc tab'
    summ['A1'].font = Font(bold=True, size=13)
    heads = [('School', None), ('Headcount', None), ('Gross', 'gross'), ('EPF (employee)', 'epf'), ('ESI employee', 'esi'), ('ESI employer', 'esiEr'), ('RRF', 'rrf'),
             ('Loan recovery', 'loanRec'), ('Net pay', 'net'), ('Bank payable', 'bank'), ('CTI', 'cti')]
    for j, (h, _) in enumerate(heads, 1):
        cc = summ.cell(3, j, h)
        cc.font = HDR_FONT
        cc.fill = HDR_FILL
        summ.column_dimensions[L(j)].width = 15
    ents = ['HES', 'LMS1', 'LMS2', 'LMS3', 'LMS4', 'LMS5', 'LMS6']
    ent_rng = "'Salary Calc'!$%s$%d:$%s$%d" % (LET['ent'], FIRST, LET['ent'], last)
    inc_rng = "'Salary Calc'!$%s$%d:$%s$%d" % (LET['incl'], FIRST, LET['incl'], last)
    for i, e in enumerate(ents, 4):
        summ.cell(i, 1, e)
        summ.cell(i, 2, '=COUNTIFS(%s,A%d,%s,"Yes")' % (ent_rng, i, inc_rng))
        for j, (h, key) in enumerate(heads[2:], 3):
            summ.cell(i, j, '=SUMIFS(\'Salary Calc\'!$%s$%d:$%s$%d,%s,$A%d,%s,"Yes")' % (LET[key], FIRST, LET[key], last, ent_rng, i, inc_rng)).number_format = '#,##0'
    tr = 4 + len(ents)
    summ.cell(tr, 1, 'TOTAL').font = BOLD
    for j in range(2, len(heads) + 1):
        cc = summ.cell(tr, j, '=SUM(%s4:%s%d)' % (L(j), L(j), tr - 1))
        cc.font = BOLD
        cc.number_format = '#,##0'

    # ---- Formula Guide -----------------------------------------------------------------------------
    guide['A1'] = 'Formula guide -- every column of Salary Calc: what it does, where its numbers come from, which backend code it mirrors'
    guide['A1'].font = Font(bold=True, size=13)
    for j, h in enumerate(['#', 'Column', 'Salary Calc col', 'What it calculates', 'Where the data comes from', 'Backend code mirrored', 'Excel formula (row %d)' % FIRST], 1):
        cc = guide.cell(3, j, h)
        cc.font = HDR_FONT
        cc.fill = HDR_FILL
    for i, cd in enumerate(COLS, 1):
        f = fill(cd['formula'], FIRST) if cd['formula'] else '(typed)'
        for j, v in enumerate([i, cd['header'], LET[cd['key']], cd['desc'], cd['src'], cd['code'], f], 1):
            cc = guide.cell(3 + i, j, v)
            cc.alignment = Alignment(wrap_text=True, vertical='top')
    for j, w in enumerate([5, 28, 8, 60, 50, 46, 90], 1):
        guide.column_dimensions[L(j)].width = w
    guide.freeze_panes = 'A4'

    # ---- Read Me ---------------------------------------------------------------------------------------
    lines = [
        ('PAYROLL PROOFREADING WORKBOOK', 'title'),
        (banner, 'warn'),
        ('', None),
        ('What this is', 'h'),
        ('Every number that makes up a month\'s salary, recalculated with live Excel formulas, one row per employee on the "Salary Calc" tab. The formulas follow the portal\'s backend code step by step (apps-script/payroll.gs and payroll-calc.gs), so a row here should always equal what the portal shows.', None),
        ('', None),
        ('How to proofread a month', 'h'),
        ('1. In Google Sheets, open each tab named below (LMCS Payroll sheet, plus PayRoll Constants and PayRoll Rates from the LMCS-Salary-PayScale sheet) and use File > Download > CSV (current sheet).', None),
        ('2. Open the CSV, copy everything below the header row, and paste it into the matching tab here starting at cell A2 (clear the old data first; leave columns U/V/W/X/Y/Q/K/M/AI/F -- the grey helper columns -- alone). Paste numbers as numbers and dates as dates.', None),
        ('3. Salary Calc tab: type the 1st of the month to check in the yellow Month cell (B4). The Settings block and every row recalculate.', None),
        ('4. Compare the "NET PAY" column with the portal / register. For a locked month the last columns do it for you ("Matches register?").', None),
        ('5. Hover any column header for the plain-English formula; hover a green/yellow cell to see the exact tab and row its value was fetched from. The "Formula Guide" tab lists all of it in one table.', None),
        ('6. A new employee: copy any row on Salary Calc down one and type the Employee Code in column A. Colour key: green = fetched from Salary Master / Rates, yellow = fetched from this month\'s inputs / ledger / loans, white = calculated, orange = results, grey = checks.', None),
        ('', None),
        ('Input tabs (paste fresh exports here)', 'h'),
        ('PayRoll Constants, PayRoll Rates  <- LMCS-Salary-PayScale sheet', None),
        ('Salary Master, Monthly Inputs, Adjustments, Ledger, Loans, Payroll Register  <- LMCS Payroll sheet', None),
        ('Tabs hold up to: Salary Master 3,000 rows · Monthly Inputs 12,000 · Adjustments 6,000 · Ledger 15,000 · Loans 1,000 · Payroll Register 15,000. Beyond that, drag the grey helper-column formulas down and widen the ranges.', None),
        ('', None),
        ('Rules worth knowing while proofreading', 'h'),
        ('EPF and ESI are both rounded UP to the next rupee (confirmed with Pawan Sir, 1 Oct 2026). ESI coverage for Apr-Sep / Oct-Mar is decided once, by the full-month wages in force on the 1st of the period (or on joining), and holds for the whole period even if pay later crosses ₹21,000.', None),
        ('Wages for EPF/ESI = Basic + ADA + DA only. Tuition, CL encashment and other earnings never enter it.', None),
        ('A person is skipped if they left before the 1st of the month (Last Working Day, else Status = Left), or their school has no PayRoll Rates row, or they have no Date of Joining.', None),
        ('If a code has several Salary Master rows (raises, transfers) the one with the latest Effective From on or before the month end is used. If two rows share the same Effective From the first one on the tab is picked (the portal prefers the later one) -- avoid duplicates.', None),
        ('Held salary (F&F / grievance) does not change Net Pay; it only moves it out of Bank payable.', None),
    ]
    for i, (t, kind) in enumerate(lines, 1):
        cc = ws_help.cell(i, 1, t)
        cc.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'title':
            cc.font = Font(bold=True, size=16)
        elif kind == 'h':
            cc.font = Font(bold=True, size=12, color='1F3864')
        elif kind == 'warn':
            cc.font = Font(bold=True, color='C00000')
    ws_help.column_dimensions['A'].width = 140
    ws_help.sheet_properties.tabColor = '1F3864'
    calc.sheet_properties.tabColor = '00B050'
    summ.sheet_properties.tabColor = '00B050'
    wb.save(path)
    return dict(first=FIRST, last=last, letters=LET, keys=KEYS, header_row=HDR_ROW)

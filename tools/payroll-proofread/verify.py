#!/usr/bin/env python3
"""Proves the proofreading workbook's formulas equal the real payroll backend.

For several random seeds it makes a fake organisation (every school, joiners, leavers, raises, holds,
releases, loans, RRF balances, adjustments incl. TDS, text/date mixes, EPF ceilings, RRF stop-at-target
on/off), builds the workbook, recalculates it in LibreOffice, runs the same data through the real
backend (tests/payroll harness) and compares every money column for every employee, over several months.

  python3 tools/payroll-proofread/verify.py [seeds]      (needs LibreOffice: soffice, and node)
"""
import datetime as dt
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
import build_workbook as B  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
SOFFICE = shutil.which('soffice') or '/opt/homebrew/bin/soffice'
D = dt.date
ENT = ['HES', 'LMS 1', 'LMS 2', 'LMS 3', 'LMS 4', 'LMS 5', 'LMS 6']
MONTHS = ['2026-09', '2026-10', '2027-03', '2026-04', '2027-04']


def fake(seed):
    rnd = random.Random(seed)
    consts = [['RRF Y1 %', 12], ['RRF Y2 %', 9], ['RRF Y3 %', 6], ['EPF Cap Salary', rnd.choice([15000, 25000])], ['EPF Rate %', 8.33],
              ['PF Rate %', 3.67], ['ESI Threshold', 21000], ['ESI Employer %', 3.25], ['ESI Employee %', 0.75]]
    if rnd.random() < .7:
        consts += [['RRF Target Months', rnd.choice([3, 2])], ['RRF Stop At Target (1=Yes)', rnd.choice([0, 1])]]
    if rnd.random() < .7:
        consts += [['Gratuity Provision %', 5], ['CL Encashment Divisor', rnd.choice([30, 26])]]
    rates = [[e, rnd.choice([5, 10]), rnd.choice([25, 35]), 'N' if e == 'LMS 6' else 'Y', 'N' if e in ('LMS 6', 'HES') and rnd.random() < .6 else 'Y'] for e in ENT]
    sm, used = [], set()
    for i in range(70):
        ent = rnd.choice(ENT) if rnd.random() < .9 else rnd.choice(ENT).replace(' ', '')
        doj = D(rnd.randint(2009, 2026), rnd.randint(1, 12), rnd.randint(1, 28))
        if rnd.random() < .12:
            doj = D(2026, 9, rnd.randint(1, 28))
        code = 'T%02d/%02d/%02d/%03d' % (i, doj.year % 100, doj.month, i)
        used.add(code)
        basic = rnd.choice([5240, 6199, 8086, 9945, 10878, 13766, 16625, 22003, 28617, 55502, rnd.randint(4500, 40000)])
        eff = max(doj, D(2026, 3, 1)) if rnd.random() < .5 else doj
        status, lwd = 'Active', ''
        x = rnd.random()
        if x < .06:
            status, lwd = 'Left', D(2026, rnd.choice([2, 7, 9, 10]), rnd.randint(1, 28))
        elif x < .10:
            status = 'Left'
        elif x < .13:
            lwd = D(rnd.choice([2026, 2027]), rnd.choice([3, 9, 10, 12]), rnd.randint(1, 28))
        epf, rrf = rnd.choice('YYYN'), rnd.choice('YYYN')
        tuition = rnd.choice([0, 0, 0, 500, 1500])
        tds = rnd.choice([0, 0, 0, 0, 2500])
        def row(effdate, b):
            ef = effdate.isoformat() if rnd.random() < .15 else effdate  # some dates arrive as text
            dj = doj.isoformat() if rnd.random() < .15 else doj
            return [code, 'Person %d' % i, ent, rnd.choice(['PRT', 'Helper', 'Peon Cum Driver']), dj, ef, b, epf, rrf, tuition, 'Bank', '', '', '', '', status, '', lwd, tds]
        sm.append(row(eff, basic))
        if rnd.random() < .3:  # a raise / transfer row
            e2 = D(rnd.choice([2026, 2027]), rnd.choice([4, 9, 10]), 1)
            if e2 > eff:
                r2 = row(e2, int(basic * rnd.choice([1.03, 1.0609, 1.5])))
                if rnd.random() < .3:
                    r2[2] = rnd.choice(ENT)
                sm.append(r2)
    codes = sorted(used)
    mi, adj, led, loans = [], [], [], []
    for mo in ['2026-09', '2026-10', '2027-03', '2026-04', '2027-04']:
        y, m = int(mo[:4]), int(mo[5:])
        for c in rnd.sample(codes, 30):
            md = D(y, m, 1)
            mi.append([md if rnd.random() < .6 else mo, c, rnd.choice(['', '', 28, 29.5, 26.25, 0, 15]), rnd.choice(['', 0, 2, 3.5]), rnd.choice(['', '', 'Y', 'N']),
                       '', '', '', '', 'x', '', rnd.choice(['', '', 'fnf', 'grievance', 'F', 'G', 'x']), rnd.choice(['', '', 1500, 4000]), rnd.choice(['', 'Y', 'N']), ''])
        for c in rnd.sample(codes, 12):
            t = rnd.choice(['Arrears', 'Travel Allowance', 'Fine', 'Other Deduction', 'TDS', 'Child Fee Recovery', 'Advance / Loan Recovery', 'Admission Incentive'])
            adj.append(['A%d' % len(adj), D(y, m, 1) if rnd.random() < .7 else mo, c, t, 'Earning' if t in ('Arrears', 'Travel Allowance', 'Admission Incentive') else 'Deduction', rnd.choice([300, 1200, 2500, 750]), '', 'x', ''])
    for c in codes:
        for acct, kind in [('RRF', 'Deduction'), ('RRF', 'Opening'), ('Security', 'Opening'), ('Held Salary', 'Deduction'), ('Loan', 'Opening')]:
            if rnd.random() < .45:
                mm = rnd.choice(['2026-03', '2026-08', '2026-09', '2026-12', '2027-02', '2027-03', '2027-04'])
                led.append([D(int(mm[:4]), int(mm[5:]), 1), D(int(mm[:4]), int(mm[5:]), 1) if rnd.random() < .6 else mm, c, acct, kind,
                            rnd.choice([1500, 12000, 60000, 3000, -2000, 250000]), '', '', 'x', ''])
    for c in rnd.sample(codes, 15):
        sm_start = rnd.choice(['', D(2026, 9, 1), D(2026, 10, 1), '2027-01', D(2027, 3, 1), D(2026, 4, 1)])
        loans.append(['L%d' % len(loans), c, 'Salary Advance', '', 20000, rnd.choice([0, 2000, 5000, 15000]), sm_start, rnd.choice(['Active', 'Active', 'Closed', '']), '', '', 'x', ''])
    return dict(data={'PayRoll Constants': consts, 'PayRoll Rates': rates, 'Salary Master': sm, 'Monthly Inputs': mi, 'Adjustments': adj,
                      'Ledger': led, 'Loans': loans, 'Payroll Register': []}, codes=codes)


def js(v):
    if isinstance(v, dt.date):
        return {'d': v.isoformat()}
    return v


FIELDS = [('paid', 'paidDays'), ('basic', 'basic'), ('adaAmt', 'ada'), ('daAmt', 'da'), ('wages', 'wages'), ('epf', 'epf'), ('esi', 'esi'), ('esiEr', 'esiEmployer'),
          ('clEnc', 'clEncashment'), ('adjE', 'otherEarnings'), ('adjD', 'otherDeductions'), ('tds', 'tds'), ('gross', 'gross'), ('cti', 'cti'), ('rrf', 'rrf'),
          ('loanRec', 'loanRecovery'), ('totDed', 'totalDeductions'), ('net', 'net'), ('bank', 'bankPayable'), ('heldF', 'heldForFnF'), ('withheld', 'withheld'), ('grat', 'gratuityProvision')]


def recalc(src, outdir):
    subprocess.run([SOFFICE, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir, src], check=True, capture_output=True, timeout=600)
    return os.path.join(outdir, os.path.basename(src))


def run(seed):
    f = fake(seed)
    tmp = tempfile.mkdtemp()
    ok = True
    # backend
    inp = {'headers': {**B.HEADERS, 'Salary Master': B.HEADERS['Salary Master']}, 'months': MONTHS,
           'data': {k: [[js(c) for c in r] for r in v] for k, v in f['data'].items()}}
    jp = os.path.join(tmp, 'in.json')
    json.dump(inp, open(jp, 'w'))
    subprocess.run(['node', os.path.join(HERE, 'dump_backend.js'), jp, os.path.join(tmp, 'out.json')], check=True)
    backend = json.load(open(os.path.join(tmp, 'out.json')))
    bad = 0
    checked = 0
    for mo in MONTHS:
        wbp = os.path.join(tmp, 'wb_%s.xlsx' % mo)
        meta = B.build(wbp, f['data'], D(int(mo[:4]), int(mo[5:]), 1), 'test')
        out = recalc(wbp, os.path.join(tmp, 'rc_' + mo)) if os.makedirs(os.path.join(tmp, 'rc_' + mo), exist_ok=True) is None else None
        ws = openpyxl.load_workbook(out, data_only=True)['Salary Calc']
        let = meta['letters']
        for r in range(meta['first'], meta['last'] + 1):
            code = ws['%s%d' % (let['code'], r)].value
            incl = ws['%s%d' % (let['incl'], r)].value
            be = backend[mo].get(code)
            if incl == 'Yes' and be is None:
                print('seed', seed, mo, code, 'sheet includes, backend skips'); bad += 1; continue
            if incl != 'Yes' and be is not None:
                print('seed', seed, mo, code, 'sheet skips (%s), backend includes' % ws['%s%d' % (let['why'], r)].value); bad += 1; continue
            if be is None:
                continue
            checked += 1
            for k, bk in FIELDS:
                v = ws['%s%d' % (let[k], r)].value
                bv = be[bk]
                if v is None or abs(float(v) - float(bv)) > 1e-6:
                    print('seed', seed, mo, code, k, 'sheet', v, 'backend', bv); bad += 1
    print('seed %d: %d employee-months compared, %d mismatches' % (seed, checked, bad))
    shutil.rmtree(tmp, ignore_errors=True)
    return bad


if __name__ == '__main__':
    seeds = [int(x) for x in sys.argv[1:]] or [1, 2, 3]
    total = sum(run(s) for s in seeds)
    print('TOTAL MISMATCHES:', total)
    sys.exit(1 if total else 0)
